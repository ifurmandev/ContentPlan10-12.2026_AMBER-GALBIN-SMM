# -*- coding: utf-8 -*-
"""Збирає TikTok-план на жовтень 2026 у .xlsx для імпорту в Google Таблиці.

Календар і заміри тягнуть дані зі «Сценаріїв» і «Свят» формулами,
тож змінюєте дату чи назву відео в одному місці — оновлюється всюди.
"""
import datetime as dt
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import content_october as C

FONT = "Lexend"
GRID = "D9D9D9"
INK = "4A4A4A"
MUTED = "8C8C8C"
HOLIDAY = "B5452B"

thin = Side(style="thin", color=GRID)
WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")


def font(size=10, color=INK, bold=False, italic=False):
    return Font(name=FONT, size=size, color=color, bold=bold, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def fit_rows(ws, first, last, pt=9):
    """Явна висота рядків під обсяг тексту — Google Таблиці при імпорті не завжди розтягують їх самі."""
    import math
    for r in range(first, last + 1):
        lines = 1
        for c in ws[r]:
            if not isinstance(c.value, str) or c.value.startswith("="):
                continue
            w = ws.column_dimensions[c.column_letter].width or 10
            per = max(1.0, w * 1.05 * 9 / pt)
            n = sum(max(1, math.ceil(len(p) / per)) for p in c.value.split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = min(409, lines * pt * 1.45 + 8)


def box(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def rubric_rules(ws, rng, col_ref):
    for name, (bg, fg) in C.RUBRIC_COLORS.items():
        key = name.split(" ", 1)[1]
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{key}",{col_ref}))'],
                             fill=fill(bg), font=Font(color=fg)))


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(9, "FFFFFF", bold=True)
        c.fill = fill("5A5A5A")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 34


def title(ws, text, sub=None, span="A1:H1"):
    ws.merge_cells(span)
    c = ws[span.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT, size=22, bold=True, color="555555")
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 40
    if sub:
        a2 = "A2:" + span.split(":")[1].rstrip("0123456789") + "2"
        ws.merge_cells(a2)
        s = ws["A2"]
        s.value = sub
        s.font = font(9, MUTED)
        s.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 44


FIRST = 5  # перший рядок даних у «Сценаріях»
LAST = FIRST + len(C.S) - 1
SC = "'Сценарії'"


def build_scenarios(wb):
    ws = wb.create_sheet("Сценарії")
    ws.sheet_view.showGridLines = False
    title(ws, "Сценарії TikTok · жовтень 2026",
          "Тут живе весь контент. Змініть дату, час або назву — календар оновиться сам. "
          "Поля [__] заповнює клієнт. Статус обирайте зі списку.", span="A1:T1")
    headers = ["№", "Дата", "День", "Час", "Рубрика", "Назва відео", "Свято / привід",
               "Суть відео", "Хук: перші 1–2 с і текст на екрані", "Розкадровка",
               "Тривалість", "Звук", "Опис і хештеги", "Закріплений коментар / заклик",
               "Мета: що дасть", "KPI-орієнтир", "Референс", "Що потрібно для зйомки",
               "Статус", "Коментар клієнта"]
    widths = [4, 8, 5, 6, 17, 24, 20, 34, 30, 44, 9, 20, 30, 30, 34, 20, 22, 28, 14, 24]
    header_row(ws, 4, headers, widths)
    for i, s in enumerate(C.S):
        r = FIRST + i
        vals = [i + 1, s["date"], s["day"], s["time"], s["rubric"], s["title"],
                f"=IFERROR(VLOOKUP(B{r},'Свята'!$A$5:$B$30,2,FALSE),\"\")",
                s["essence"], s["hook"], s["board"], s["dur"], s["sound"], s["caption"],
                s["pinned"], s["goal"], s["kpi"], s["ref"], s["need"], "Ідея", ""]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = font(9)
            c.alignment = WRAP_TOP
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=r, column=2).number_format = "dd.mm"
        ws.cell(row=r, column=6).font = font(9, bold=True)
        ws.cell(row=r, column=7).font = font(8, HOLIDAY, bold=True)
    rubric_rules(ws, f"E{FIRST}:F{LAST}", f"$E{FIRST}")
    dv = DataValidation(type="list", formula1="='Рубрики'!$B$5:$B$8", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E{FIRST}:E{LAST + 20}")
    st = DataValidation(type="list",
                        formula1='"Ідея,Сценарій погоджено,Знято,Змонтовано,Опубліковано,Перенесено"',
                        allow_blank=True)
    ws.add_data_validation(st)
    st.add(f"S{FIRST}:S{LAST + 20}")
    done = ws.conditional_formatting
    done.add(f"S{FIRST}:S{LAST + 20}", FormulaRule(formula=[f'$S{FIRST}="Опубліковано"'],
                                                   fill=fill("D9EAD3"), font=Font(color="274E13")))
    fit_rows(ws, FIRST, LAST)
    ws.freeze_panes = "G5"
    return ws


def build_calendar(wb):
    ws = wb.active
    ws.title = "Календар"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 24
    ws.column_dimensions["I"].width = 2

    ws.row_dimensions[1].height = 10
    ws.merge_cells("B2:E2")
    ws["B2"] = "Жовтень"
    ws["B2"].font = Font(name=FONT, size=54, bold=True, color="555555")
    ws["B2"].alignment = Alignment(vertical="bottom")
    ws.merge_cells("F2:H2")
    ws["F2"] = "2026"
    ws["F2"].font = Font(name=FONT, size=44, bold=True, color="BDBDBD")
    ws["F2"].alignment = Alignment(horizontal="right", vertical="bottom")
    ws.row_dimensions[2].height = 78
    ws.merge_cells("B3:H3")
    ws["B3"] = "TikTok · Amber Galbin · 4 відео на тиждень · Пн 09:00 · Ср 21:00 · Пт 21:00 · Нд 11:00"
    ws["B3"].font = font(10, MUTED)
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 8

    days = ["ПОНЕДІЛОК", "ВІВТОРОК", "СЕРЕДА", "ЧЕТВЕР", "П’ЯТНИЦЯ", "СУБОТА", "НЕДІЛЯ"]
    for i, d in enumerate(days):
        c = ws.cell(row=5, column=2 + i, value=d)
        c.font = font(9, "6B6B6B")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[5].height = 28

    start = dt.date(2026, 9, 28)
    row = 6
    for week in range(5):
        dr, hr, cr = row, row + 1, row + 2
        ws.row_dimensions[dr].height = 34
        ws.row_dimensions[hr].height = 26
        ws.row_dimensions[cr].height = 72
        for i in range(7):
            col = 2 + i
            L = get_column_letter(col)
            day = start + dt.timedelta(days=week * 7 + i)
            in_month = day.month == 10
            dc = ws.cell(row=dr, column=col, value=day if in_month else None)
            dc.number_format = "d"
            dc.font = Font(name=FONT, size=20, color="5A5A5A")
            dc.alignment = Alignment(horizontal="left", vertical="top", indent=0)
            dc.border = Border(left=thin, right=thin, top=thin)

            hc = ws.cell(row=hr, column=col,
                         value=f"=IFERROR(VLOOKUP({L}{dr},'Свята'!$A$5:$B$30,2,FALSE),\"\")")
            hc.font = font(8, HOLIDAY, bold=True)
            hc.alignment = WRAP_TOP
            hc.border = Border(left=thin, right=thin)

            m = f"MATCH({L}{dr},{SC}!$B${FIRST}:$B${LAST + 20},0)"
            cc = ws.cell(row=cr, column=col, value=(
                f"=IFERROR(INDEX({SC}!$E${FIRST}:$E${LAST + 20},{m})&\"  ·  \"&"
                f"INDEX({SC}!$D${FIRST}:$D${LAST + 20},{m})&CHAR(10)&"
                f"INDEX({SC}!$F${FIRST}:$F${LAST + 20},{m}),\"\")"))
            cc.font = font(8, INK, bold=True)
            cc.alignment = WRAP_TOP
            cc.border = Border(left=thin, right=thin, bottom=thin)
        row += 3
    last_grid = row - 1
    rubric_rules(ws, f"B6:H{last_grid}", "B6")

    row += 1
    ws.cell(row=row, column=2, value="РУБРИКИ").font = font(8, MUTED, bold=True)
    for i, r in enumerate(C.RUBRICS):
        bg, fg = C.RUBRIC_COLORS[r["name"]]
        c = ws.cell(row=row, column=3 + i, value=f"{r['name']}\n{r['slot']}")
        c.font = font(8, fg, bold=True)
        c.fill = fill(bg)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c = ws.cell(row=row, column=7, value="Червоним — свята і\nінфоприводи (вкладка «Свята»)")
    c.font = font(8, HOLIDAY, bold=True)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32
    row += 2
    ws.merge_cells(f"B{row}:H{row}")
    n = ws.cell(row=row, column=2, value=(
        "Календар заповнюється автоматично з вкладки «Сценарії» (рубрика, час, назва) і «Свята». "
        "Щоб перенести відео — змініть дату в «Сценаріях». Деталі кожного відео — там само."))
    n.font = font(8, MUTED, italic=True)
    n.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[row].height = 28
    return ws


def build_rubrics(wb):
    ws = wb.create_sheet("Рубрики")
    ws.sheet_view.showGridLines = False
    title(ws, "Рубрики і правила", C.BASIS)
    headers = ["", "Рубрика", "Слот", "Суть", "Референс і цифри", "Мета: що дає", "Інші KPI", "Ціль переглядів"]
    widths = [3, 20, 11, 40, 40, 38, 24, 12]
    header_row(ws, 4, headers, widths)
    for i, r in enumerate(C.RUBRICS):
        rr = 5 + i
        bg, fg = C.RUBRIC_COLORS[r["name"]]
        vals = [i + 1, r["name"], r["slot"], r["essence"], r["ref"], r["goal"], r["kpi"], r["target"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=rr, column=j, value=v)
            c.font = font(9)
            c.alignment = WRAP_TOP
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=rr, column=2).fill = fill(bg)
        ws.cell(row=rr, column=2).font = font(9, fg, bold=True)
        ws.cell(row=rr, column=8).number_format = "# ##0"
    fit_rows(ws, 5, 4 + len(C.RUBRICS))
    rr = 5 + len(C.RUBRICS) + 1
    ws.cell(row=rr, column=2, value="Правила для всіх відео").font = font(12, "555555", bold=True)
    for k, rule in enumerate(C.RULES, start=1):
        ws.cell(row=rr + k, column=1, value=k).font = font(9, MUTED)
        ws.merge_cells(start_row=rr + k, start_column=2, end_row=rr + k, end_column=7)
        c = ws.cell(row=rr + k, column=2, value=rule)
        c.font = font(9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr + k].height = 26
    return ws


def build_holidays(wb):
    ws = wb.create_sheet("Свята")
    ws.sheet_view.showGridLines = False
    title(ws, "Свята та інфоприводи",
          "Дати перевірені на 2026 рік. Календар підтягує назву свята за датою з колонки A — "
          "додайте рядок, і свято з’явиться в календарі.", span="A1:E1")
    header_row(ws, 4, ["Дата", "Свято / подія", "Тип", "Як використати в контенті", "Відео в плані"],
               [10, 40, 18, 56, 12])
    for i, (d, name, kind, use, vid) in enumerate(C.HOLIDAYS):
        r = 5 + i
        for j, v in enumerate([d, name, kind, use, vid], start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = font(9)
            c.alignment = WRAP_TOP
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=r, column=1).number_format = "dd.mm.yyyy"
        ws.cell(row=r, column=2).font = font(9, HOLIDAY, bold=True)
    fit_rows(ws, 5, 4 + len(C.HOLIDAYS))
    return ws


def build_references(wb):
    ws = wb.create_sheet("Референси")
    ws.sheet_view.showGridLines = False
    title(ws, "Референси: що спрацювало і чому",
          "Цифри зняті з екранних записів TikTok у папці referens (округлені, як показує TikTok). "
          "ER = (вподобання + коментарі + збереження + поширення) / перегляди.", span="A1:M1")
    headers = ["Відео", "Опубліковано", "Перегляди", "Вподобання", "Коментарі", "Збереження",
               "Поширення", "ER", "Поширення / перегляди", "Що в кадрі", "Чому спрацювало / ні",
               "Що беремо в план", "Рубрика"]
    header_row(ws, 4, headers, [26, 11, 11, 11, 10, 11, 11, 8, 11, 40, 40, 36, 17])
    for i, rf in enumerate(C.REFERENCES):
        r = 5 + i
        vals = [rf["name"], rf["date"], rf["views"], rf["likes"], rf["comments"], rf["saves"],
                rf["shares"],
                f'=IF(AND(C{r}>0,COUNT(D{r}:G{r})=4),(D{r}+E{r}+F{r}+G{r})/C{r},"")',
                f'=IF(AND(C{r}>0,G{r}<>""),G{r}/C{r},"")',
                rf["frame"], rf["why"], rf["take"], rf["rubric"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = font(9)
            c.alignment = WRAP_TOP
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=r, column=1).font = font(9, bold=True)
        ws.cell(row=r, column=2).number_format = "dd.mm.yyyy"
        for j in range(3, 8):
            ws.cell(row=r, column=j).number_format = "# ##0"
        ws.cell(row=r, column=8).number_format = "0.0%"
        ws.cell(row=r, column=9).number_format = "0.00%"
    rubric_rules(ws, f"M5:M{4 + len(C.REFERENCES)}", "$M5")
    fit_rows(ws, 5, 4 + len(C.REFERENCES))
    ws.freeze_panes = "B5"
    return ws


def build_metrics(wb):
    ws = wb.create_sheet("Заміри")
    ws.sheet_view.showGridLines = False
    title(ws, "Заміри після публікації",
          "Заповнюємо через 24–48 год з TikTok Studio (білі клітинки). Сірі — рахуються самі. "
          "Порівнюйте з референсами: «Орел» — ER 6,9%, «Сирець» — 2,7%, «Доміно» — 3,8%.", span="A1:R1")
    headers = ["№", "Дата", "Відео", "Рубрика", "Перегляди", "Сер. час перегляду, %",
               "До кінця, %", "Вподобання", "Коментарі", "Збереження", "Поширення",
               "Нові підписники", "Перегляди профілю", "ER", "Поширення / перегляди",
               "Ціль переглядів", "Ціль виконано?", "Висновок: повторюємо / змінюємо"]
    header_row(ws, 4, headers, [4, 7, 30, 17, 10, 10, 9, 10, 9, 10, 9, 10, 10, 7, 10, 10, 10, 30])
    calc = fill("F3F3F3")
    for i in range(len(C.S)):
        r = FIRST + i
        src = FIRST + i
        vals = [f"={SC}!A{src}", f"={SC}!B{src}", f"={SC}!F{src}", f"={SC}!E{src}"] + [None] * 9 + [
            f'=IF(E{r}>0,(H{r}+I{r}+J{r}+K{r})/E{r},"")',
            f'=IF(E{r}>0,K{r}/E{r},"")',
            f"=IFERROR(VLOOKUP(D{r},'Рубрики'!$B$5:$H$8,7,FALSE),\"\")",
            f'=IF(E{r}="","",IF(E{r}>=P{r},"✅ так","❌ ні"))', None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = font(9)
            c.alignment = WRAP_TOP
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if isinstance(v, str) and v.startswith("="):
                c.fill = calc
        ws.cell(row=r, column=2).number_format = "dd.mm"
        for j in (5, 8, 9, 10, 11, 12, 13, 16):
            ws.cell(row=r, column=j).number_format = "# ##0"
        ws.cell(row=r, column=14).number_format = "0.0%"
        ws.cell(row=r, column=15).number_format = "0.00%"
    last = FIRST + len(C.S) - 1
    rubric_rules(ws, f"D{FIRST}:D{last}", f"$D{FIRST}")

    r0 = last + 2
    ws.cell(row=r0, column=3, value="Підсумок по рубриках").font = font(12, "555555", bold=True)
    sub = ["Рубрика", "Опубліковано", "Сер. перегляди", "Сер. ER", "Сер. поширення"]
    for j, h in enumerate(sub):
        c = ws.cell(row=r0 + 1, column=3 + j, value=h)
        c.font = font(9, "FFFFFF", bold=True)
        c.fill = fill("5A5A5A")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    rng = lambda col: f"${col}${FIRST}:${col}${last}"
    for k in range(len(C.RUBRICS)):
        rr = r0 + 2 + k
        vals = [f"='Рубрики'!B{5 + k}",
                f'=COUNTIFS({rng("D")},C{rr},{rng("E")},">0")',
                f'=IFERROR(AVERAGEIFS({rng("E")},{rng("D")},C{rr},{rng("E")},">0"),"")',
                f'=IFERROR(AVERAGEIFS({rng("N")},{rng("D")},C{rr},{rng("E")},">0"),"")',
                f'=IFERROR(AVERAGEIFS({rng("O")},{rng("D")},C{rr},{rng("E")},">0"),"")']
        for j, v in enumerate(vals):
            c = ws.cell(row=rr, column=3 + j, value=v)
            c.font = font(9)
            c.fill = calc
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=rr, column=5).number_format = "# ##0"
        ws.cell(row=rr, column=6).number_format = "0.0%"
        ws.cell(row=rr, column=7).number_format = "0.00%"
    rubric_rules(ws, f"C{r0 + 2}:C{r0 + 1 + len(C.RUBRICS)}", f"$C{r0 + 2}")
    ws.freeze_panes = "D5"
    return ws


def main(out):
    wb = Workbook()
    build_calendar(wb)
    build_scenarios(wb)
    build_rubrics(wb)
    build_holidays(wb)
    build_references(wb)
    build_metrics(wb)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1 if ws.title == "Календар" else 0
        ws.print_options.horizontalCentered = True
    tmp = out + ".tmp"
    wb.save(tmp)
    # Перепаковуємо з максимальним стисненням і без docProps — файл менший для завантаження.
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for item in zin.infolist():
            if item.filename.startswith("docProps/"):
                continue
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b'<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>', b"")
                data = data.replace(b'<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>', b"")
            if item.filename == "_rels/.rels":
                import re
                data = re.sub(rb'<Relationship[^>]*docProps[^>]*/>', b"", data)
            zout.writestr(item.filename, data)
    import os
    os.remove(tmp)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Amber_Galbin_TikTok_October_2026.xlsx")
